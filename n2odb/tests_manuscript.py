"""Phase-A tests for manuscript ingestion (manuscript plan §10).

No network and no live model call. Covers: text extraction (PDF, DOCX-with-table,
pasted text, broken PDF), deterministic parsers (DOI + decimal/DMS coordinates),
and the §5 seam (stub returns skipped; fixture fake returns the full contract) end
to end through the upload/extract endpoints.
"""
from __future__ import annotations

import io
import json
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from n2odb.tests_common import AuthedWorkbenchTestCase
from django.urls import reverse

from n2odb.models_upload import ManuscriptSource, UploadSubmission
from n2odb.services import doi_enrich, manuscript_extract, manuscript_regex
from n2odb.services.manuscript_fixtures import fake_extract_facts
from n2odb.services.manuscript_merge import derive_metadata, merge_metadata
from n2odb.services.manuscript_text import extract_text


# --- fixtures ----------------------------------------------------------------

def _make_pdf(text: str) -> bytes:
    """Assemble a minimal single-page PDF containing ``text``, with a correct
    cross-reference table (byte offsets computed as we go)."""
    def esc(s: str) -> str:
        return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    stream = f"BT /F1 12 Tf 72 700 Td ({esc(text)}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"

    xref_pos = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += (
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF"
        % (len(objects) + 1, xref_pos)
    )
    return bytes(out)


def _make_docx_with_table() -> bytes:
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("Study of nitrous oxide emissions.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Sand"
    table.cell(0, 1).text = "46%"
    table.cell(1, 0).text = "Clay"
    table.cell(1, 1).text = "24%"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- text extraction ---------------------------------------------------------

class ExtractTextTests(TestCase):
    def test_pdf_extracts_text_and_page_count(self):
        pdf = _make_pdf("Hello world DOI 10.1234/abcd")
        result = extract_text("pdf", pdf, filename="paper.pdf")
        self.assertEqual(result.error, "")
        self.assertEqual(result.page_count, 1)
        self.assertIn("Hello world", result.text)
        self.assertEqual(result.char_count, len(result.text))

    def test_docx_includes_table_cells(self):
        result = extract_text("docx", _make_docx_with_table(), filename="s.docx")
        self.assertEqual(result.error, "")
        self.assertIsNone(result.page_count)
        self.assertIn("nitrous oxide", result.text)
        self.assertIn("Sand", result.text)
        self.assertIn("46%", result.text)

    def test_pasted_text_dehyphenates(self):
        result = extract_text("text", "soil mois-\nture content")
        self.assertEqual(result.error, "")
        self.assertIn("soil moisture content", result.text)

    def test_paragraph_breaks_preserved(self):
        result = extract_text("text", "Line one\nwrapped.\n\nNew paragraph.")
        self.assertIn("Line one wrapped.", result.text)
        self.assertIn("\n\nNew paragraph.", result.text)

    def test_cluster_rows_groups_lines_left_to_right(self):
        from n2odb.services.manuscript_text import _cluster_rows

        # Two visual lines; words given out of order, second line's baseline
        # slightly jittered within tolerance.
        words = [
            {"text": "world", "top": 10.0, "x0": 60},
            {"text": "Hello", "top": 10.0, "x0": 10},
            {"text": "second", "top": 30.5, "x0": 40},
            {"text": "line", "top": 29.0, "x0": 70},
        ]
        rows = _cluster_rows(words)
        joined = [" ".join(w["text"] for w in row) for row in rows]
        self.assertEqual(joined, ["Hello world", "second line"])

    def test_broken_pdf_fails_without_raising(self):
        result = extract_text("pdf", b"%PDF-1.4 not really a pdf", filename="x.pdf")
        self.assertNotEqual(result.error, "")
        self.assertEqual(result.text, "")
        self.assertEqual(result.char_count, 0)

    @override_settings(WORKBENCH_DOC_MAX_MB=0)
    def test_oversized_document_degrades(self):
        result = extract_text("pdf", b"%PDF-1.4 anything", filename="x.pdf")
        self.assertNotEqual(result.error, "")
        self.assertEqual(result.text, "")


# --- deterministic parsers ---------------------------------------------------

class RegexParserTests(TestCase):
    def test_doi(self):
        entry = manuscript_regex.find_doi("Available at https://doi.org/10.1234/abc.def-5.")
        self.assertEqual(entry["value"], "10.1234/abc.def-5")
        self.assertEqual(entry["source"], "regex")
        self.assertIn("10.1234", entry["source_quote"])

    def test_decimal_lat_long_order(self):
        found = manuscript_regex.find_coordinates(
            "The site is located at -33.87, 151.21 in New South Wales."
        )
        self.assertAlmostEqual(found["site_latitude"]["value"], -33.87, places=2)
        self.assertAlmostEqual(found["site_longitude"]["value"], 151.21, places=2)

    def test_dms_southern_western_hemispheres(self):
        found = manuscript_regex.find_coordinates(
            "Coordinates 33°52'11\"S, 151°12'30\"W were recorded."
        )
        self.assertAlmostEqual(found["site_latitude"]["value"], -33.8697, places=3)
        self.assertAlmostEqual(found["site_longitude"]["value"], -151.2083, places=3)

    def test_dms_ring_above_degree_glyph(self):
        # Real journal PDFs typeset the degree mark as ˚ (U+02DA) with a prime
        # (U+2032) for minutes, e.g. "(42˚24′ N, 85˚24′ W, and 288 masl)".
        found = manuscript_regex.find_coordinates(
            "located in southwest Michigan (42˚24′ N, 85˚24′ W, and 288 masl)"
        )
        self.assertAlmostEqual(found["site_latitude"]["value"], 42.4, places=3)
        self.assertAlmostEqual(found["site_longitude"]["value"], -85.4, places=3)

    def test_degree_decimal_hemispheres(self):
        found = manuscript_regex.find_coordinates(
            "The KBS site is at 42.44°N, 85.38°W in Michigan."
        )
        self.assertAlmostEqual(found["site_latitude"]["value"], 42.44, places=2)
        self.assertAlmostEqual(found["site_longitude"]["value"], -85.38, places=2)

    def test_coordinate_format_variations(self):
        # A wide spread of real-world formats should all resolve to ~(42.4, -85.4).
        cases = [
            "42°24'30\"N, 85°24'30\"W",                        # DMS with seconds
            "42°24.5′N, 85°24.5′W",                            # decimal minutes
            "N 42°24', W 85°24'",                              # hemisphere first
            "42.4°N, 85.4°W",                                  # decimal + degree
            "42.4˚ N, 85.4˚ W",                                # ring-above glyph
            "latitude 42.4, longitude -85.4",                  # axis-labelled decimals
            "lat. 42°24′ N; long. 85°24′ W",                   # axis-labelled DMS
            "coordinates 42.4°, -85.4°",                       # signed degree pair + cue
            "situated at 42.4, -85.4",                         # unmarked pair + cue
        ]
        for text in cases:
            with self.subTest(text=text):
                found = manuscript_regex.find_coordinates(text)
                self.assertIn("site_latitude", found, text)
                self.assertIn("site_longitude", found, text)
                self.assertGreater(found["site_latitude"]["value"], 0, text)
                self.assertLess(found["site_longitude"]["value"], 0, text)
                # latitude magnitude in the low 40s, longitude in the mid 80s
                self.assertAlmostEqual(found["site_latitude"]["value"], 42.4, delta=0.2)
                self.assertAlmostEqual(found["site_longitude"]["value"], -85.4, delta=0.2)

    def test_european_decimal_comma(self):
        # A global corpus: "42,4" is 42.4 when a degree/label disambiguates it.
        for text in [
            "The site lies at 42,4°N, 85,4°W in the US corn belt.",
            "latitude 42,4; longitude -85,4",
            "42°24,5′N, 85°24,5′W",
        ]:
            with self.subTest(text=text):
                found = manuscript_regex.find_coordinates(text)
                self.assertAlmostEqual(found["site_latitude"]["value"], 42.4, delta=0.2)
                self.assertAlmostEqual(found["site_longitude"]["value"], -85.4, delta=0.2)

    def test_us_comma_separator_not_treated_as_decimal(self):
        # Bare comma-separated integers must NOT be read as one European number,
        # even next to a cue — there is no degree/hemisphere to disambiguate.
        self.assertEqual(
            manuscript_regex.find_coordinates("coordinates for 42, 85 and 90 plots"),
            {},
        )

    def test_label_word_long_is_not_a_false_coordinate(self):
        # "long-term" / "long ... years" must not be read as a longitude.
        found = manuscript_regex.find_coordinates(
            "long-term cropping systems observed for 85 years"
        )
        self.assertEqual(found, {})

    def test_nitrogen_is_not_a_coordinate(self):
        # A bare hemisphere letter after a decimal is nitrogen here, not a coord.
        found = manuscript_regex.find_coordinates(
            "Urea was applied at 5.2 N ha-1 and again at 120.0 N over the season."
        )
        self.assertEqual(found, {})

    def test_unanchored_decimal_pair_is_ignored(self):
        # Two adjacent decimals with no coordinate cue must not be guessed as coords.
        found = manuscript_regex.find_coordinates(
            "Soil organic carbon ranged from 12.5, 43.2 g kg-1 across plots."
        )
        self.assertEqual(found, {})

    def test_parse_all_keys(self):
        meta = manuscript_regex.parse_all(
            "DOI 10.9999/xyz at 40.0°N, 105.0°W."
        )
        self.assertIn("pub_doi", meta)
        self.assertIn("site_latitude", meta)
        self.assertGreater(meta["site_latitude"]["value"], 0)
        self.assertLess(meta["site_longitude"]["value"], 0)


_ABSTRACT_BODY = (
    "Nitrous oxide (N2O) emissions were measured from two agroecosystems using "
    "static chambers and the flux-gradient technique over two growing seasons. "
    "High spatial variability was observed and long pulses followed fertilisation "
    "and irrigation events at the site in eastern Washington."
)


class FindAbstractTests(TestCase):
    def test_heading_delimited_slice(self):
        text = (
            "Title line\n\nAbstract\n\n" + _ABSTRACT_BODY
            + "\n\nKeywords: N2O, maize\n\n1. Introduction\n\nThe rest."
        )
        entry = manuscript_regex.find_abstract(text)
        self.assertIsNotNone(entry)
        self.assertIn("Nitrous oxide", entry["value"])
        self.assertNotIn("Introduction", entry["value"])
        self.assertNotIn("Keywords", entry["value"])
        self.assertEqual(entry["source"], "regex")

    def test_same_line_heading_with_colon(self):
        text = "Abstract: " + _ABSTRACT_BODY + "\n\nIntroduction\n\nBody."
        entry = manuscript_regex.find_abstract(text)
        self.assertIsNotNone(entry)
        self.assertTrue(entry["value"].startswith("Nitrous oxide"))

    def test_short_fragment_is_rejected(self):
        text = "Abstract\n\nToo short.\n\nIntroduction\n\nBody."
        self.assertIsNone(manuscript_regex.find_abstract(text))

    def test_missing_abstract_returns_none(self):
        self.assertIsNone(
            manuscript_regex.find_abstract("Introduction\n\nNo abstract here.")
        )

    def test_parse_all_includes_abstract(self):
        text = (
            "DOI 10.1234/abc.\n\nAbstract\n\n" + _ABSTRACT_BODY
            + "\n\nIntroduction\n\nRest."
        )
        meta = manuscript_regex.parse_all(text)
        self.assertIn("pub_abstract", meta)
        self.assertIn("pub_doi", meta)


class DeriveMetadataTests(TestCase):
    def test_doi_yields_a_link(self):
        out = derive_metadata(
            {"pub_doi": {"value": "10.1029/2019JG005032", "source": "regex"}}
        )
        self.assertEqual(out["pub_link"]["value"], "https://doi.org/10.1029/2019jg005032")
        self.assertEqual(out["pub_link"]["source"], "derived")

    def test_existing_link_is_left_alone(self):
        out = derive_metadata({
            "pub_doi": {"value": "10.1/x"},
            "pub_link": {"value": "https://example.org/paper", "source": "document"},
        })
        self.assertEqual(out["pub_link"]["value"], "https://example.org/paper")

    def test_no_doi_no_link(self):
        out = derive_metadata({"pub_title": {"value": "x"}})
        self.assertNotIn("pub_link", out)

    def test_does_not_mutate_input(self):
        src = {"pub_doi": {"value": "10.1/x"}}
        derive_metadata(src)
        self.assertNotIn("pub_link", src)

    def test_coordinates_become_signed_decimals(self):
        out = derive_metadata({
            "site_latitude": {"value": "26°58′16.8″ S", "unit": "°"},
            "site_longitude": {"value": "85°24′", "unit": "°"},  # paper's W dropped
        })
        self.assertAlmostEqual(out["site_latitude"]["value"], -26.971333)
        self.assertEqual(out["site_latitude"]["original_value"], "26°58′16.8″ S")
        self.assertEqual(out["site_longitude"]["value"], 85.4)
        self.assertIn("check the sign", out["site_longitude"]["note"])

    def test_impossible_dms_is_left_as_written(self):
        out = derive_metadata({"site_longitude": {"value": "151°82′85.3″ E", "unit": "°"}})
        self.assertEqual(out["site_longitude"]["value"], "151°82′85.3″ E")

    def test_soil_and_climate_units_convert_to_canonical(self):
        out = derive_metadata({
            "sand": {"value": "365", "unit": "g kg−1"},
            "bulk_density": {"value": "1450", "unit": "kg m-3"},
            "map": {"value": "89.0", "unit": "cm"},
            "mat": {"value": "50", "unit": "°F"},
            "clay": {"value": "33 (0–10 cm), 31 (10–20 cm)", "unit": "%"},  # profile: untouched
        })
        self.assertEqual((out["sand"]["value"], out["sand"]["unit"]), ("36.5", "%"))
        self.assertEqual(out["bulk_density"]["value"], "1.45")
        self.assertEqual(out["map"]["value"], "890")
        self.assertEqual((out["mat"]["value"], out["mat"]["unit"]), ("10", "°C"))
        self.assertEqual(out["clay"]["value"], "33 (0–10 cm), 31 (10–20 cm)")


class DoiEnrichServiceTests(TestCase):
    def _fake_urlopen(self, payload, status=200):
        return io.BytesIO(json.dumps(payload).encode())

    def test_returns_cleaned_abstract(self):
        payload = {"abstract": "<jats:p>Abstract: Nitrous oxide fluxes were measured.</jats:p>"}
        with mock.patch("urllib.request.urlopen", return_value=self._fake_urlopen(payload)):
            self.assertEqual(
                doi_enrich.fetch_abstract("10.1234/abc"),
                "Nitrous oxide fluxes were measured.",
            )

    def test_missing_abstract_returns_none(self):
        with mock.patch("urllib.request.urlopen", return_value=self._fake_urlopen({})):
            self.assertIsNone(doi_enrich.fetch_abstract("10.1234/abc"))

    def test_network_error_returns_none(self):
        with mock.patch("urllib.request.urlopen", side_effect=OSError("boom")):
            self.assertIsNone(doi_enrich.fetch_abstract("10.1234/abc"))

    def test_blank_doi_returns_none(self):
        self.assertIsNone(doi_enrich.fetch_abstract(""))

    def test_markup_only_cleans_to_none(self):
        self.assertIsNone(doi_enrich.clean_abstract("<jats:p></jats:p>"))

    def test_fetch_citation_uses_crossref(self):
        payload = {"status": "ok", "message": {
            "title": ["N2O emissions from two agroecosystems"],
            "author": [{"family": "Waldo", "given": "Sarah"},
                       {"family": "Russell", "given": "Eric S."}],
            "issued": {"date-parts": [[2019]]},
            "container-title": ["Journal of Geophysical Research: Biogeosciences"],
            "volume": "124", "page": "1887-1904", "DOI": "10.1029/2019JG005032",
        }}
        with mock.patch("urllib.request.urlopen", return_value=self._fake_urlopen(payload)):
            result = doi_enrich.fetch_citation("10.1029/2019JG005032")
        self.assertEqual(result[1], "crossref")
        self.assertEqual(
            result[0],
            "Waldo, S., & Russell, E. (2019). N2O emissions from two agroecosystems. "
            "Journal of Geophysical Research: Biogeosciences, 124, 1887-1904. "
            "https://doi.org/10.1029/2019JG005032",
        )

    def test_fetch_citation_falls_back_to_semantic_scholar(self):
        crossref = {"status": "ok", "message": {}}  # no title -> not usable
        s2 = {
            "title": "A dataset paper",
            "authors": [{"name": "Ada Lovelace"}],
            "year": 2020,
            "venue": "Data Journal",
            "externalIds": {"DOI": "10.1/abc"},
        }
        with mock.patch("urllib.request.urlopen", side_effect=[
            self._fake_urlopen(crossref), self._fake_urlopen(s2),
        ]):
            result = doi_enrich.fetch_citation("10.1/abc")
        self.assertEqual(result[1], "semantic_scholar")
        self.assertEqual(
            result[0],
            "Lovelace, A. (2020). A dataset paper. Data Journal. "
            "https://doi.org/10.1/abc",
        )

    def test_fetch_citation_network_error_returns_none(self):
        with mock.patch("urllib.request.urlopen", side_effect=OSError("boom")):
            self.assertIsNone(doi_enrich.fetch_citation("10.1/abc"))

    def test_fetch_citation_blank_doi_returns_none(self):
        self.assertIsNone(doi_enrich.fetch_citation(""))


class MergeTests(TestCase):
    def test_prefers_deterministic_value_keeps_module_quote(self):
        # Values disagree well beyond tolerance → flagged; deterministic wins,
        # module's quote is retained.
        regex_meta = {"site_latitude": {"value": -33.87, "source_quote": "", "source": "regex"}}
        module_meta = {"site_latitude": {"value": -30.0, "source_quote": "around 30 S", "confidence": 0.7}}
        merged = merge_metadata(regex_meta, module_meta)
        self.assertEqual(merged["site_latitude"]["value"], -33.87)
        self.assertEqual(merged["site_latitude"]["source_quote"], "around 30 S")
        self.assertIn("mismatch", merged["site_latitude"])

    def test_close_values_do_not_flag_mismatch(self):
        # A rounding-level difference (within tolerance) is not a mismatch.
        merged = merge_metadata(
            {"site_latitude": {"value": -33.87, "source": "regex"}},
            {"site_latitude": {"value": -33.9, "source_quote": "33.9 S"}},
        )
        self.assertNotIn("mismatch", merged["site_latitude"])

    def test_union_of_fields(self):
        merged = merge_metadata(
            {"pub_doi": {"value": "10.1/x", "source": "regex"}},
            {"primary_crop": {"value": "maize"}},
        )
        self.assertIn("pub_doi", merged)
        self.assertIn("primary_crop", merged)

    def test_abstract_prefers_the_module_value(self):
        # Unlike every other field, the module's verbatim abstract wins over the
        # heading slice; the regex entry is only a fallback.
        merged = merge_metadata(
            {"pub_abstract": {"value": "regex slice", "source": "regex"}},
            {"pub_abstract": {"value": "module abstract", "source": "document"}},
        )
        self.assertEqual(merged["pub_abstract"]["value"], "module abstract")
        self.assertEqual(merged["pub_abstract"]["source"], "document")

    def test_abstract_falls_back_to_regex_when_module_empty(self):
        merged = merge_metadata(
            {"pub_abstract": {"value": "regex slice", "source": "regex"}},
            {"pub_abstract": {"value": None, "source": "document"}},
        )
        self.assertEqual(merged["pub_abstract"]["value"], "regex slice")


# --- the §5 seam -------------------------------------------------------------

class ExtractionSeamTests(TestCase):
    @override_settings(WORKBENCH_LLM_ENABLED=False)
    def test_stub_returns_skipped_when_disabled(self):
        result = manuscript_extract.extract_facts("some text", [])
        self.assertEqual(result["status"], "skipped")
        self.assertEqual(result["metadata"], {})
        self.assertEqual(result["events"], [])

    def test_fixture_fake_returns_full_contract(self):
        facts = fake_extract_facts("text", [])
        self.assertEqual(facts["status"], "ok")
        self.assertIn("primary_crop", facts["metadata"])
        self.assertEqual(len(facts["events"]), 1)
        self.assertEqual(facts["events"][0]["type"], "input_application")

    def test_fixture_fake_respects_already_have(self):
        facts = fake_extract_facts("text", ["primary_crop"])
        self.assertNotIn("primary_crop", facts["metadata"])


# --- endpoints (end to end, no model) ----------------------------------------

@override_settings(WORKBENCH_ENRICH_ENABLED=False)  # never reach out to the network here
class UploadEndpointTests(AuthedWorkbenchTestCase):
    def setUp(self):
        super().setUp()
        self.submission = UploadSubmission.objects.create()
        self.token = self.submission.public_token

    def test_paste_text_creates_source_and_regex_suggestions(self):
        url = reverse("workbench-document", args=[self.token])
        resp = self.client.post(
            url,
            {"text": "Field study (DOI 10.1234/soil.2018) at latitude -33.87, 151.21."},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data["status"], "extracted")
        self.assertGreaterEqual(data["regex_suggestions"], 2)

        doc = ManuscriptSource.objects.get(pk=data["document_id"])
        self.assertEqual(doc.kind, "text")
        self.assertEqual(doc.extraction_status, "extracted")

        self.submission.refresh_from_db()
        meta = self.submission.manuscript_facts["metadata"]
        self.assertIn("pub_doi", meta)
        self.assertIn("site_latitude", meta)

    @override_settings(WORKBENCH_MAX_PASTED_TEXT_CHARS=10)
    def test_reject_oversize_pasted_text(self):
        resp = self.client.post(
            reverse("workbench-document", args=[self.token]),
            {"text": "x" * 11},
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("too long", resp.json()["error"].lower())

    def test_reject_unsupported_file(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        url = reverse("workbench-document", args=[self.token])
        bad = SimpleUploadedFile("note.txt", b"just text", content_type="text/plain")
        resp = self.client.post(url, {"document": bad})
        self.assertEqual(resp.status_code, 400)

    @override_settings(WORKBENCH_MAX_UPLOAD_MB=0)  # 0 MB → any non-empty file is too large
    def test_reject_oversize_file(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        url = reverse("workbench-document", args=[self.token])
        big = SimpleUploadedFile("p.pdf", b"%PDF-1.4\n" + b"x" * 200,
                                 content_type="application/pdf")
        resp = self.client.post(url, {"document": big})
        self.assertEqual(resp.status_code, 400)
        self.assertIn("too large", resp.json()["error"].lower())

    @override_settings(WORKBENCH_MAX_FILES_PER_SUBMISSION=1)
    def test_reject_over_file_cap(self):
        url = reverse("workbench-document", args=[self.token])
        self.assertEqual(self.client.post(url, {"text": "first"}).status_code, 200)
        resp = self.client.post(url, {"text": "second"})
        self.assertEqual(resp.status_code, 400)
        self.assertIn("files", resp.json()["error"].lower())

    @override_settings(WORKBENCH_MAX_FILES_PER_SUBMISSION=1)
    def test_document_and_data_files_share_file_cap(self):
        self.assertEqual(
            self.client.post(
                reverse("workbench-document", args=[self.token]),
                {"text": "first"},
            ).status_code,
            200,
        )
        from django.core.files.uploadedfile import SimpleUploadedFile

        data = SimpleUploadedFile("data.csv", b"date,n2o_flux\n2020-01-01,1.2\n")
        resp = self.client.post(
            reverse("workbench-data", args=[self.token]),
            {"data": data},
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("files", resp.json()["error"].lower())

    @override_settings(WORKBENCH_LLM_ENABLED=False)
    def test_extract_endpoint_skipped_when_llm_disabled(self):
        # Capture text first.
        self.client.post(
            reverse("workbench-document", args=[self.token]),
            {"text": "Methods. DOI 10.1234/abc at 40.0 N, 105.0 W."},
        )
        resp = self.client.post(reverse("workbench-document-extract", args=[self.token]))
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        # Module disabled → skipped, but deterministic DOI/coords still present.
        self.assertEqual(data["status"], "skipped")
        self.assertIn("pub_doi", data["manuscript_facts"]["metadata"])

    def test_extract_endpoint_with_fixture_fake(self):
        self.client.post(
            reverse("workbench-document", args=[self.token]),
            {"text": "Methods paragraph mentioning maize. DOI 10.1234/abc."},
        )
        with mock.patch.object(
            manuscript_extract, "extract_facts", side_effect=fake_extract_facts
        ):
            resp = self.client.post(
                reverse("workbench-document-extract", args=[self.token])
            )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data["status"], "ok")
        meta = data["manuscript_facts"]["metadata"]
        # Fixture metadata plus the deterministic DOI, and the event carried through.
        self.assertIn("primary_crop", meta)
        self.assertIn("pub_doi", meta)
        self.assertEqual(data["events_count"], 1)

        self.submission.refresh_from_db()
        self.assertIsNotNone(self.submission.manuscript_extracted_at)

    def test_submit_data_page_renders_stage0(self):
        resp = self.client.get(reverse("submit-data"))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode()
        self.assertIn('id="dropzone"', body)
        self.assertIn("csrfmiddlewaretoken", body)
        # The three endpoint URLs are wired into the page for the JS to call.
        self.assertIn("/document/", body)

    def test_document_text_endpoint_serves_extracted_text(self):
        self.client.post(
            reverse("workbench-document", args=[self.token]),
            {"text": "Readable body text for the panel."},
        )
        resp = self.client.get(reverse("workbench-document-text", args=[self.token]))
        self.assertEqual(resp.status_code, 200)
        self.assertIn(b"Readable body text", resp.content)


@override_settings(WORKBENCH_ENRICH_ENABLED=False)
class ExtractionCostGuardTests(AuthedWorkbenchTestCase):
    """Dedupe + per-account daily cap on the (billable) extract endpoint."""

    def setUp(self):
        super().setUp()
        self.submission = UploadSubmission.objects.create(owner=self.user)
        self.token = self.submission.public_token
        # Attach a document (pasted text gets a sha256 like any source).
        self.client.post(
            reverse("workbench-document", args=[self.token]),
            {"text": "Methods paragraph mentioning maize. DOI 10.1234/abc."},
        )
        self.extract_url = reverse("workbench-document-extract", args=[self.token])

    def _extract(self, **post):
        with mock.patch.object(
            manuscript_extract, "extract_facts", side_effect=fake_extract_facts
        ) as spy:
            resp = self.client.post(self.extract_url, post)
        return resp, spy

    def test_first_extract_records_one_call_and_stores_sha(self):
        resp, spy = self._extract()
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(resp.json()["deduped"])
        self.assertEqual(spy.call_count, 1)
        from n2odb.models_upload import LlmExtractionCall
        self.assertEqual(LlmExtractionCall.objects.filter(user=self.user).count(), 1)
        self.submission.refresh_from_db()
        self.assertTrue(self.submission.manuscript_source_sha256)

    def test_rerun_unchanged_is_deduped_no_new_call(self):
        self._extract()  # first, successful
        resp, spy = self._extract()  # second, unchanged, no force
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.json()["deduped"])
        self.assertEqual(spy.call_count, 0)  # model was never called
        from n2odb.models_upload import LlmExtractionCall
        self.assertEqual(LlmExtractionCall.objects.filter(user=self.user).count(), 1)

    def test_force_reruns_even_when_unchanged(self):
        self._extract()
        resp, spy = self._extract(force="1")
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(resp.json()["deduped"])
        self.assertEqual(spy.call_count, 1)
        from n2odb.models_upload import LlmExtractionCall
        self.assertEqual(LlmExtractionCall.objects.filter(user=self.user).count(), 2)


@override_settings(WORKBENCH_ENRICH_ENABLED=False)
class ExtractionDailyCapTests(TestCase):
    """The cap only bites real (assist-enabled) calls and exempts staff."""

    def setUp(self):
        super().setUp()
        self.user = get_user_model().objects.create_user(
            username="contrib", email="c@university.edu", password="pw",  # non-staff
        )
        self.client.force_login(self.user)
        self.submission = UploadSubmission.objects.create(owner=self.user)
        self.token = self.submission.public_token
        self.client.post(
            reverse("workbench-document", args=[self.token]),
            {"text": "Methods paragraph mentioning maize. DOI 10.1234/abc."},
        )
        self.extract_url = reverse("workbench-document-extract", args=[self.token])

    def _force_extract(self):
        # force=1 bypasses dedupe so each call is a real, counted extraction.
        with mock.patch("n2odb.views_workbench.llm_assist.is_enabled", return_value=True), \
             mock.patch.object(manuscript_extract, "extract_facts",
                               side_effect=fake_extract_facts):
            return self.client.post(self.extract_url, {"force": "1"})

    @override_settings(WORKBENCH_LLM_DAILY_CAP=2)
    def test_cap_blocks_after_limit(self):
        self.assertEqual(self._force_extract().status_code, 200)
        self.assertEqual(self._force_extract().status_code, 200)
        blocked = self._force_extract()
        self.assertEqual(blocked.status_code, 429)
        self.assertIn("limit", blocked.json()["error"].lower())

    @override_settings(WORKBENCH_LLM_DAILY_CAP=1)
    def test_staff_are_exempt(self):
        self.user.is_staff = True
        self.user.save(update_fields=["is_staff"])
        self.assertEqual(self._force_extract().status_code, 200)
        self.assertEqual(self._force_extract().status_code, 200)  # over cap, still ok

    @override_settings(WORKBENCH_LLM_DAILY_CAP=0)
    def test_cap_of_zero_disables_the_ceiling(self):
        for _ in range(3):
            self.assertEqual(self._force_extract().status_code, 200)


@override_settings(WORKBENCH_LLM_ENABLED=False, WORKBENCH_ENRICH_ENABLED=False)
class EnrichmentEndpointTests(AuthedWorkbenchTestCase):
    """The DOI→abstract lookup: the gated automatic path and the manual retry.

    The language-model pass is disabled so the suite stays fully offline; the
    enrichment lookups are mocked per test.
    """

    def setUp(self):
        super().setUp()
        self.submission = UploadSubmission.objects.create(owner=self.user)
        self.token = self.submission.public_token
        self.client.post(
            reverse("workbench-document", args=[self.token]),
            {"text": "Methods paragraph. DOI 10.1234/abc."},
        )

    @override_settings(WORKBENCH_ENRICH_ENABLED=True)
    def test_auto_enrich_fills_a_missing_abstract_and_citation(self):
        with mock.patch.object(
            doi_enrich, "fetch_abstract", return_value="An abstract from Semantic Scholar."
        ), mock.patch.object(
            doi_enrich, "fetch_citation",
            return_value=("Smith, J. (2019). A paper. Journal, 1, 2. https://doi.org/10.1234/abc",
                          "crossref"),
        ):
            resp = self.client.post(
                reverse("workbench-document-extract", args=[self.token])
            )
        self.assertEqual(resp.status_code, 200)
        facts = resp.json()["manuscript_facts"]
        self.assertEqual(
            facts["metadata"]["pub_abstract"]["value"], "An abstract from Semantic Scholar."
        )
        self.assertEqual(facts["metadata"]["pub_abstract"]["source"], "semantic_scholar")
        self.assertEqual(facts["enrichment"]["abstract"], "filled")
        self.assertEqual(facts["metadata"]["pub_citation"]["source"], "crossref")
        self.assertEqual(facts["enrichment"]["citation"], "filled")

    def test_auto_enrich_off_makes_no_network_call(self):
        with mock.patch.object(doi_enrich, "fetch_abstract") as abs_spy, \
             mock.patch.object(doi_enrich, "fetch_citation") as cite_spy:
            resp = self.client.post(
                reverse("workbench-document-extract", args=[self.token])
            )
        self.assertEqual(resp.status_code, 200)
        abs_spy.assert_not_called()
        cite_spy.assert_not_called()

    def test_manual_enrich_fills_the_abstract(self):
        with mock.patch.object(
            doi_enrich, "fetch_abstract", return_value="Manual abstract."
        ), mock.patch.object(doi_enrich, "fetch_citation", return_value=None):
            resp = self.client.post(
                reverse("workbench-document-enrich", args=[self.token])
            )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()["status"], "ok")
        self.submission.refresh_from_db()
        self.assertEqual(
            self.submission.manuscript_facts["metadata"]["pub_abstract"]["value"],
            "Manual abstract.",
        )

    def test_manual_enrich_fills_the_citation(self):
        cite = "Smith, J. (2019). A paper. Journal, 1, 2. https://doi.org/10.1234/abc"
        with mock.patch.object(doi_enrich, "fetch_abstract", return_value=None), \
             mock.patch.object(doi_enrich, "fetch_citation", return_value=(cite, "crossref")):
            resp = self.client.post(
                reverse("workbench-document-enrich", args=[self.token])
            )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()["status"], "ok")
        self.assertEqual(resp.json()["citation"], cite)
        self.submission.refresh_from_db()
        self.assertEqual(
            self.submission.manuscript_facts["metadata"]["pub_citation"]["value"], cite
        )
        self.assertEqual(
            self.submission.manuscript_facts["metadata"]["pub_citation"]["source"], "crossref"
        )

    def test_manual_enrich_without_a_doi_explains(self):
        sub = UploadSubmission.objects.create(owner=self.user)
        resp = self.client.post(
            reverse("workbench-document-enrich", args=[sub.public_token])
        )
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.json()["status"], "no_doi")
