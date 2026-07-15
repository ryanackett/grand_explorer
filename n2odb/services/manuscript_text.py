"""Plain-text extraction from an uploaded source document.

Pure and deterministic — no network, no model calls. Dispatch on ``kind`` and
return normalized, human-readable text plus a couple of counts. Parsing failures
never raise into the request: they return an ``ExtractedText`` with empty text and
an ``error`` set, and the caller records ``extraction_status="failed"``.

See manuscript-extraction-implementation-plan §3.

Licensing note: we deliberately use pdfplumber (MIT), pypdf (BSD) and python-docx
(MIT). Do NOT add PyMuPDF/fitz — it is AGPL and this project is heading toward an
open-source release.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass

from django.conf import settings


@dataclass
class ExtractedText:
    text: str
    page_count: int | None
    char_count: int
    # Set when parsing failed; the caller maps this to extraction_status="failed".
    error: str = ""


class DocumentTooLarge(Exception):
    """Raised internally when a document exceeds the configured guards."""


def _limit_mb() -> int:
    return int(getattr(settings, "WORKBENCH_DOC_MAX_MB", 32))


def _limit_pages() -> int:
    return int(getattr(settings, "WORKBENCH_DOC_MAX_PAGES", 60))


# --- normalization -----------------------------------------------------------

_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_DEHYPHEN_RE = re.compile(r"(\w)-\n(\w)")
# A single newline not adjacent to another newline = an intra-paragraph line
# wrap; a blank line (two+ newlines) is a real paragraph break we keep.
_INTRA_PARA_RE = re.compile(r"(?<!\n)\n(?!\n)")
_MULTI_BLANK_RE = re.compile(r"\n{3,}")
_HSPACE_RE = re.compile(r"[ \t]+")


def _normalize(text: str) -> str:
    """Make raw extracted text readable and stable.

    De-hyphenate line-break splits, fold intra-paragraph newlines to spaces while
    keeping blank-line paragraph breaks, collapse whitespace, strip control chars,
    and NFC-normalize. Output must stay human-readable for the document panel.
    """
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _DEHYPHEN_RE.sub(r"\1\2", text)
    text = _INTRA_PARA_RE.sub(" ", text)
    text = _MULTI_BLANK_RE.sub("\n\n", text)
    text = _CONTROL_RE.sub("", text)
    # Collapse runs of spaces/tabs and trim trailing space on each line.
    text = _HSPACE_RE.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = unicodedata.normalize("NFC", text)
    return text.strip()


# --- per-kind parsers --------------------------------------------------------

# A PDF that parses but yields almost nothing is probably scanned or column-heavy;
# below this many chars we retry with the other library before giving up.
_NEAR_EMPTY_CHARS = 40


def _extract_pdf(data: bytes) -> tuple[str, int | None]:
    """Try pdfplumber first (better multi-column/layout), fall back to pypdf.

    Raises if *both* backends fail to read the file at all (a corrupt PDF), so the
    caller records it as failed. A file that parses but yields little/no text (a
    scanned/empty PDF) is not an error — it returns empty text with a page count.
    """
    text, pages, primary_exc = "", None, None
    try:
        text, pages = _pdf_with_pdfplumber(data)
    except DocumentTooLarge:
        raise
    except Exception as exc:  # noqa: BLE001
        primary_exc = exc

    if len((text or "").strip()) < _NEAR_EMPTY_CHARS:
        try:
            alt_text, alt_pages = _pdf_with_pypdf(data)
        except DocumentTooLarge:
            raise
        except Exception as exc:  # noqa: BLE001
            # Both backends failed and we have nothing: the file is unreadable.
            if primary_exc is not None and not (text or "").strip() and pages is None:
                raise primary_exc
        else:
            if pages is None:
                pages = alt_pages
            if len((alt_text or "").strip()) > len((text or "").strip()):
                text, pages = alt_text, alt_pages
    return text, pages


# Tuning for the target corpus: two-column agronomy journal PDFs.
#
# * ``x_tolerance_ratio`` sets the word-break gap as a fraction of each glyph's
#   font size, instead of pdfplumber's fixed 3pt default. The default merges
#   adjacent words in the small, justified type of narrow columns
#   ("croppingsystems,whichchallenges…"); a font-relative gap recovers the
#   spaces while still adapting to larger title/heading type. 0.15 was tuned
#   against real papers in this format.
# * Rotated glyphs (``upright=False``) are the publisher's margin watermark
#   ("Downloaded from …/ See the Terms and Conditions …") stamped down the side
#   of every page; dropping them keeps that boilerplate out of the body text.
_X_TOLERANCE_RATIO = 0.15
# Words whose baselines are within this many points count as the same visual line.
_LINE_Y_TOLERANCE = 3


def _cluster_rows(words: list[dict]) -> list[list[dict]]:
    """Group extracted words into visual lines (top-to-bottom, left-to-right)."""
    rows: list[list] = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if rows and abs(w["top"] - rows[-1][0]) <= _LINE_Y_TOLERANCE:
            rows[-1][1].append(w)
        else:
            rows.append([w["top"], [w]])
    return [sorted(ws, key=lambda w: w["x0"]) for _, ws in rows]


def _pdfplumber_page_text(page) -> str:
    words = page.extract_words(
        x_tolerance_ratio=_X_TOLERANCE_RATIO,
        use_text_flow=False,
        extra_attrs=["upright"],
    )
    words = [w for w in words if w.get("upright", True)]
    if not words:
        # Rare (e.g. a page whose text is entirely rotated); fall back to the
        # plain extractor so we don't silently drop the page.
        return page.extract_text(x_tolerance_ratio=_X_TOLERANCE_RATIO) or ""
    return "\n".join(
        " ".join(w["text"] for w in row) for row in _cluster_rows(words)
    )


def _pdf_with_pdfplumber(data: bytes) -> tuple[str, int | None]:
    import io

    import pdfplumber

    parts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        page_count = len(pdf.pages)
        if page_count > _limit_pages():
            raise DocumentTooLarge(
                f"{page_count} pages exceeds the {_limit_pages()}-page limit"
            )
        for page in pdf.pages:
            parts.append(_pdfplumber_page_text(page))
    return "\n\n".join(parts), page_count


def _pdf_with_pypdf(data: bytes) -> tuple[str, int | None]:
    import io

    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    page_count = len(reader.pages)
    if page_count > _limit_pages():
        raise DocumentTooLarge(
            f"{page_count} pages exceeds the {_limit_pages()}-page limit"
        )
    parts = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(parts), page_count


def _extract_docx(data: bytes) -> tuple[str, int | None]:
    """Join paragraph text and walk tables (site characteristics are often tabular)."""
    import io

    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            # Deduplicate horizontally-merged cells that repeat the same text.
            deduped = []
            for cell in cells:
                if not deduped or deduped[-1] != cell:
                    deduped.append(cell)
            line = " | ".join(c for c in deduped if c)
            if line:
                parts.append(line)
    return "\n\n".join(parts), None


# --- public entry point ------------------------------------------------------


def extract_text(kind: str, data: bytes | str, *, filename: str = "") -> ExtractedText:
    """Extract normalized plain text from a document. Never raises.

    ``kind`` is one of "pdf", "docx", "text". ``data`` is raw bytes for pdf/docx
    and a str (or bytes) for pasted text. On any failure the returned
    ``ExtractedText`` has empty text and a non-empty ``error``.
    """
    try:
        if kind == "text":
            raw = data.decode("utf-8", "replace") if isinstance(data, bytes) else (data or "")
            text, pages = raw, None
        else:
            if not isinstance(data, (bytes, bytearray)):
                raise ValueError(f"{kind} requires bytes, got {type(data).__name__}")
            data = bytes(data)
            max_bytes = _limit_mb() * 1024 * 1024
            if len(data) > max_bytes:
                raise DocumentTooLarge(
                    f"{len(data)} bytes exceeds the {_limit_mb()} MB limit"
                )
            if kind == "pdf":
                text, pages = _extract_pdf(data)
            elif kind == "docx":
                text, pages = _extract_docx(data)
            else:
                raise ValueError(f"Unsupported document kind: {kind!r}")

        normalized = _normalize(text)
        return ExtractedText(
            text=normalized, page_count=pages, char_count=len(normalized)
        )
    except DocumentTooLarge as exc:
        return ExtractedText(text="", page_count=None, char_count=0, error=str(exc))
    except Exception as exc:  # noqa: BLE001 — degrade gracefully, never 500
        return ExtractedText(
            text="",
            page_count=None,
            char_count=0,
            error=f"Could not read document: {exc}",
        )
